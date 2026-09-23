#!/usr/bin/env python3
"""Export the canonical CSVs back to node-list.xlsx / edge-list.xlsx for editing.

    python3 scripts/to_xlsx.py                  # overwrite the workbooks in place
    python3 scripts/to_xlsx.py --out review/    # write copies elsewhere
    python3 scripts/to_xlsx.py --year 1660 --month 2

Frozen header row, an autofilter, and sensible column widths, so a month can be
read down the screen. Coordinates keep 6 decimal places (~0.1 m) rather than
Excel's default float rendering.
"""
import argparse
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import pepys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    raise SystemExit("needs openpyxl:  pip install openpyxl")

WIDTHS = {"id": 26, "name": 30, "location": 30, "latitude": 12, "longitude": 12,
          "notes": 60, "source": 26, "target": 26, "year": 9, "month": 7,
          "day": 6, "mode": 8, "companions": 40}
NUMERIC = {"month", "day", "latitude", "longitude"}


def build(path, cols, rows, title):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(cols)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDD6C7")
        cell.alignment = Alignment(vertical="center")

    for r in rows:
        out = []
        for c in cols:
            v = (r.get(c) or "").strip()
            if v and c in NUMERIC:
                try:
                    v = int(v) if c in ("month", "day") else float(v)
                except ValueError:
                    pass
            out.append(v)
        ws.append(out)

    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(c, 18)
        if c in ("latitude", "longitude"):
            for cell in ws[get_column_letter(i)][1:]:
                cell.number_format = "0.000000"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)
    print("%-22s %4d rows" % (os.path.basename(path), len(rows)))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=pepys.REPO, help="output directory")
    ap.add_argument("--year", help="export only this year's edges")
    ap.add_argument("--month", type=int, help="export only this month (needs --year)")
    args = ap.parse_args()
    os.makedirs(args.out, exist_ok=True)

    nodes = pepys.read_nodes()
    build(os.path.join(args.out, "node-list.xlsx"), pepys.NODE_COLS, nodes, "nodes")

    edges = []
    for y in ([args.year] if args.year else pepys.edge_years()):
        edges.extend(pepys.read_edges(y))
    if args.month:
        edges = [e for e in edges if e["month"].strip() == str(args.month)]
    build(os.path.join(args.out, "edge-list.xlsx"), pepys.EDGE_COLS, edges, "edges")


if __name__ == "__main__":
    main()
